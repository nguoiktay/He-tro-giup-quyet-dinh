import os
import sys
import logging
import re
import openpyxl
from datetime import datetime
import bcrypt

thuMucGoc = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if thuMucGoc not in sys.path:
    sys.path.insert(0, thuMucGoc)

from database.connection import layPhienKetNoi, dongPhienKetNoi, taoDongCo
from database.models import (
    TaiKhoan,
    HocSinh,
    DiemMonHoc,
    SoThichHocSinh,
    TruongDaiHoc,
    NganhHoc,
    PhuongAn,
    ToHopXetTuyen,
    DiemChuan,
    NguonDuLieu,
    MauDanhGia,
    PhienBanDuLieu,
    DanhSachNguyenVong,
    NguyenVong,
)
from config import CauHinh, CauHinhKiemThu

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")


def xacDinhNhomNganh(tenNganh):
    ten = tenNganh.lower()
    if any(k in ten for k in ["máy tính", "công nghệ thông tin", "phần mềm", "dữ liệu", "trí tuệ nhân tạo", "an toàn thông tin", "mạng", "hệ thống thông tin"]):
        return "CongNgheThongTin"
    if any(k in ten for k in ["kinh tế", "quản trị", "kinh doanh", "tài chính", "ngân hàng", "kế toán", "kiểm toán", "marketing", "thương mại"]):
        return "KinhTeQuanTri"
    if any(k in ten for k in ["kỹ thuật", "cơ khí", "điện", "điện tử", "tự động hóa", "xây dựng", "giao thông", "kiến trúc"]):
        return "KyThuatCongNghe"
    if any(k in ten for k in ["y", "dược", "điều dưỡng", "răng", "hàm", "mặt", "y tế", "sức khỏe"]):
        return "YDuoc"
    if any(k in ten for k in ["sư phạm", "giáo dục", "ngôn ngữ", "văn học", "lịch sử", "địa lý", "triết học", "báo chí", "truyền thông"]):
        return "SuPhamNhanVan"
    if any(k in ten for k in ["luật", "ngoại giao", "quốc tế", "chính trị", "quản lý nhà nước", "an ninh", "cảnh sát"]):
        return "LuatNgoaiGiao"
    if any(k in ten for k in ["nông nghiệp", "lâm nghiệp", "thủy sản", "chăn nuôi", "thú y", "sinh học"]):
        return "NongLamThuySan"
    return "Khac"


def chuanHoaKhoiTruong(khoi):
    if not khoi:
        return "CongLap"
    khoiStr = str(khoi).lower()
    if "quân đội" in khoiStr or "công an" in khoiStr:
        return "QuanDoiCongAn"
    if "ngoài lập" in khoiStr or "tư thục" in khoiStr or "dân lập" in khoiStr:
        return "NgoaiLap"
    return "CongLap"


def chuanHoaVungMien(vung):
    if not vung:
        return "MienBac"
    vungStr = str(vung).lower()
    if "trung" in vungStr:
        return "MienTrung"
    if "nam" in vungStr:
        return "MienNam"
    return "MienBac"


def uocLuongHocPhiKy(khoiTruong):
    if khoiTruong == "QuanDoiCongAn":
        return 0.0
    if khoiTruong == "NgoaiLap":
        return 25000000.0
    return 12500000.0


def napDuLieuExcel(duongDanTep, gioiHanBanGhi=2500):
    phien = layPhienKetNoi()
    try:
        logging.info("Bắt đầu đọc tệp Excel: %s", duongDanTep)
        wb = openpyxl.load_workbook(duongDanTep, read_only=True)

        nguon = phien.query(NguonDuLieu).filter_by(tenNguon="Bộ GD&ĐT Tuyển sinh 2021-2026").first()
        if not nguon:
            nguon = NguonDuLieu(
                tenNguon="Bộ GD&ĐT Tuyển sinh 2021-2026",
                loaiNguon="ExcelTongHop",
                duongDan=duongDanTep,
                moTa="Dữ liệu tổng hợp điểm chuẩn, chỉ tiêu và tổ hợp tuyển sinh đại học toàn quốc giai đoạn 2021 - 2026."
            )
            phien.add(nguon)
            phien.flush()

        sheetToHop = wb["05_DanhMuc_280_ToHopMon"]
        cacToHopDaCo = set(r[0] for r in phien.query(ToHopXetTuyen.maToHop).all())
        toHopMoi = []
        for i, row in enumerate(sheetToHop.iter_rows(min_row=2, values_only=True)):
            if not row or not row[1]:
                continue
            maToHop = str(row[1]).strip()
            if not re.match(r"^[A-Za-z0-9]+$", maToHop):
                continue
            if maToHop not in cacToHopDaCo:
                tenToHop = str(row[2]).strip() if row[2] else maToHop
                danhSachMon = str(row[3]).strip() if row[3] else tenToHop
                toHopMoi.append(ToHopXetTuyen(
                    maToHop=maToHop,
                    tenToHop=tenToHop[:100],
                    danhSachMon=danhSachMon[:100]
                ))
                cacToHopDaCo.add(maToHop)
        if toHopMoi:
            phien.bulk_save_objects(toHopMoi)
            phien.flush()
            logging.info("Đã nạp %d tổ hợp môn xét tuyển.", len(toHopMoi))

        toHopCoBan = [
            ("A00", "Toán, Vật lý, Hóa học", "Toan, VatLy, HoaHoc"),
            ("A01", "Toán, Vật lý, Tiếng Anh", "Toan, VatLy, NgoaiNgu"),
            ("B00", "Toán, Hóa học, Sinh học", "Toan, HoaHoc, SinhHoc"),
            ("C00", "Ngữ văn, Lịch sử, Địa lý", "Van, LichSu, DiaLy"),
            ("D01", "Toán, Ngữ văn, Tiếng Anh", "Toan, Van, NgoaiNgu"),
            ("D07", "Toán, Hóa học, Tiếng Anh", "Toan, HoaHoc, NgoaiNgu"),
        ]
        for ma, ten, ds in toHopCoBan:
            if ma not in cacToHopDaCo:
                phien.add(ToHopXetTuyen(maToHop=ma, tenToHop=ten, danhSachMon=ds))
                cacToHopDaCo.add(ma)
        phien.flush()

        sheetTruong = wb["03_Danh_Sach_587_Truong_DH"]
        cacTruongDaCo = set(r[0] for r in phien.query(TruongDaiHoc.maTruong).all())
        truongMoi = []
        for row in sheetTruong.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            maTruong = str(row[1]).strip().upper()
            if maTruong not in cacTruongDaCo:
                tenTruong = str(row[2]).strip() if row[2] else maTruong
                khoiTruong = chuanHoaKhoiTruong(row[3])
                vungMien = chuanHoaVungMien(row[5])
                tinhThanh = str(row[6]).strip() if row[6] else "Toàn quốc"
                truongMoi.append(TruongDaiHoc(
                    maTruong=maTruong[:30],
                    tenTruong=tenTruong[:255],
                    khoiTruong=khoiTruong,
                    tinhThanh=tinhThanh[:100],
                    vungMien=vungMien,
                    website=None
                ))
                cacTruongDaCo.add(maTruong)
        if truongMoi:
            phien.bulk_save_objects(truongMoi)
            phien.flush()
            logging.info("Đã nạp %d trường đại học.", len(truongMoi))

        sheetDiem = wb["01_Ma_Tran_Diem_Chuan_Cac_Nam"]
        cacNganhDaCo = set(r[0] for r in phien.query(NganhHoc.maNganh).all())
        cacPhuongAnDaCo = set(r[0] for r in phien.query(PhuongAn.maPhuongAn).all())
        cacDiemChuanDaCo = set(
            (r[0], r[1], r[2]) for r in phien.query(DiemChuan.maPhuongAn, DiemChuan.maToHop, DiemChuan.nam).all()
        )

        nganhMoi = []
        phuongAnMoi = []
        diemChuanMoi = []

        demBanGhi = 0
        for row in sheetDiem.iter_rows(min_row=2, values_only=True):
            if not row or not row[1] or not row[6]:
                continue
            maTruong = str(row[1]).strip().upper()
            if maTruong not in cacTruongDaCo:
                continue

            tenNganh = str(row[6]).strip()
            nhomNganh = xacDinhNhomNganh(tenNganh)
            maNganhTao = f"N{abs(hash(tenNganh)) % 1000000:06d}"
            if maNganhTao not in cacNganhDaCo:
                nganhMoi.append(NganhHoc(
                    maNganh=maNganhTao,
                    tenNganh=tenNganh[:255],
                    nhomNganh=nhomNganh,
                    moTa=f"Chương trình đào tạo ngành {tenNganh}"
                ))
                cacNganhDaCo.add(maNganhTao)

            phuongThuc = "DiemThiTHPT"
            if row[7] and "học bạ" in str(row[7]).lower():
                phuongThuc = "HocBa"
            elif row[7] and "đánh giá" in str(row[7]).lower():
                phuongThuc = "DanhGiaNangLuc"

            maPhuongAn = f"{maTruong}_{maNganhTao}_{phuongThuc}"[:60]
            if maPhuongAn not in cacPhuongAnDaCo:
                khoiTruong = "CongLap"
                hocPhiUocLuong = uocLuongHocPhiKy(khoiTruong)
                phuongAnMoi.append(PhuongAn(
                    maPhuongAn=maPhuongAn,
                    maTruong=maTruong,
                    maNganh=maNganhTao,
                    tenChuongTrinh=tenNganh[:255],
                    coSo="Cơ sở chính",
                    phuongThuc=phuongThuc,
                    hocPhiKy=hocPhiUocLuong,
                    donViTien="VND",
                    namTuyenSinh=2026,
                    chiTieu=100,
                    dieuKienXetTuyen="Tốt nghiệp THPT và đạt ngưỡng đảm bảo chất lượng đầu vào."
                ))
                cacPhuongAnDaCo.add(maPhuongAn)

            toHopStr = str(row[8]).strip() if row[8] else "A00"
            dsToHop = [t.strip() for t in re.split(r"[;,/]", toHopStr) if t.strip()]
            if not dsToHop:
                dsToHop = ["A00"]

            cotNam = [
                (2021, row[9]),
                (2022, row[10]),
                (2023, row[11]),
                (2024, row[12]),
                (2025, row[13]),
                (2026, row[14]),
            ]
            for nam, diemVal in cotNam:
                if diemVal is not None:
                    try:
                        diemFloat = float(diemVal)
                        if 10.0 <= diemFloat <= 30.0:
                            for th in dsToHop:
                                if th in cacToHopDaCo and (maPhuongAn, th, nam) not in cacDiemChuanDaCo:
                                    diemChuanMoi.append(DiemChuan(
                                        maPhuongAn=maPhuongAn,
                                        maToHop=th,
                                        nam=nam,
                                        diemTrungTuyen=diemFloat,
                                        thangDiem=30.0,
                                        nguonDuLieu="DeAnTuyenSinh",
                                        ghiChu=None
                                    ))
                                    cacDiemChuanDaCo.add((maPhuongAn, th, nam))
                    except (ValueError, TypeError):
                        pass

            demBanGhi += 1
            if gioiHanBanGhi and demBanGhi >= gioiHanBanGhi:
                break

        if nganhMoi:
            phien.bulk_save_objects(nganhMoi)
            phien.flush()
            logging.info("Đã nạp %d ngành học mới.", len(nganhMoi))

        if phuongAnMoi:
            phien.bulk_save_objects(phuongAnMoi)
            phien.flush()
            logging.info("Đã nạp %d phương án tuyển sinh.", len(phuongAnMoi))

        if diemChuanMoi:
            phien.bulk_save_objects(diemChuanMoi)
            phien.flush()
            logging.info("Đã nạp %d bản ghi điểm chuẩn lịch sử.", len(diemChuanMoi))

        wb.close()
        phien.commit()
        logging.info("Hoàn tất nạp dữ liệu danh mục từ Excel!")
    except Exception as e:
        phien.rollback()
        logging.error("Lỗi khi nạp dữ liệu Excel: %s", e)
        raise
    finally:
        dongPhienKetNoi()


def taoTaiKhoanMacDinhVaMauDanhGia():
    phien = layPhienKetNoi()
    try:
        emailAdmin = "admin@dss.edu.vn"
        admin = phien.query(TaiKhoan).filter_by(email=emailAdmin).first()
        if not admin:
            matKhauAdminHash = bcrypt.hashpw(b"AdminPassword@2026", bcrypt.gensalt()).decode("utf-8")
            admin = TaiKhoan(
                email=emailAdmin,
                matKhauHash=matKhauAdminHash,
                vaiTro="quanTriVien",
                trangThai="kichHoat"
            )
            phien.add(admin)
            phien.flush()
            logging.info("Đã tạo tài khoản quản trị mặc định: %s", emailAdmin)

        emailHocSinh = "hocsinh@dss.edu.vn"
        tkHs = phien.query(TaiKhoan).filter_by(email=emailHocSinh).first()
        if not tkHs:
            matKhauHsHash = bcrypt.hashpw(b"StudentPassword@2026", bcrypt.gensalt()).decode("utf-8")
            tkHs = TaiKhoan(
                email=emailHocSinh,
                matKhauHash=matKhauHsHash,
                vaiTro="hocSinh",
                trangThai="kichHoat"
            )
            phien.add(tkHs)
            phien.flush()

            hs = HocSinh(
                maTaiKhoan=tkHs.maTaiKhoan,
                hoTen="Nguyễn Văn An",
                soDienThoai="0912345678",
                namDuTuyen=2026,
                nganSachToiDa=25000000.0,
                kyTinhNganSach="nam",
                diaBanUuTien="MienBac",
                batBuocNganSach=False,
                batBuocDiaBan=False,
                mucTieuNgheNghiep="Trở thành kỹ sư trí tuệ nhân tạo và phát triển phần mềm",
                uuTienCaNhan="Ưu tiên trường đào tạo công nghệ uy tín tại Hà Nội"
            )
            phien.add(hs)
            phien.flush()

            diemMacDinh = [
                ("Toan", 8.6),
                ("VatLy", 8.4),
                ("HoaHoc", 7.8),
                ("NgoaiNgu", 8.2),
                ("Van", 7.5),
                ("SinhHoc", 6.5),
                ("LichSu", 7.0),
                ("DiaLy", 7.2),
                ("GDCD", 8.5),
            ]
            for mon, diem in diemMacDinh:
                phien.add(DiemMonHoc(
                    maHocSinh=hs.maHocSinh,
                    tenMon=mon,
                    diemSo=diem,
                    thangDiem=10.0,
                    loaiDiem="thpt",
                    namHocKy="2026"
                ))

            phien.add(SoThichHocSinh(maHocSinh=hs.maHocSinh, nhomNganh="CongNgheThongTin", mucDoThich=5))
            phien.add(SoThichHocSinh(maHocSinh=hs.maHocSinh, nhomNganh="KyThuatCongNghe", mucDoThich=4))
            phien.add(SoThichHocSinh(maHocSinh=hs.maHocSinh, nhomNganh="KinhTeQuanTri", mucDoThich=3))

            dsnv = DanhSachNguyenVong(
                maHocSinh=hs.maHocSinh,
                tenDanhSach="Danh sách nguyện vọng chính thức",
                phienBan=1
            )
            phien.add(dsnv)
            phien.flush()

            logging.info("Đã tạo tài khoản học sinh mẫu: %s", emailHocSinh)

        danhSachNhomNganh = [
            "CongNgheThongTin", "KinhTeQuanTri", "KyThuatCongNghe",
            "YDuoc", "SuPhamNhanVan", "LuatNgoaiGiao", "NongLamThuySan"
        ]
        danhSachVung = ["MienBac", "MienTrung", "MienNam"]
        
        cacHocSinhId = [hs.maHocSinh]
        matKhauChungHash = bcrypt.hashpw(b"StudentPassword@2026", bcrypt.gensalt()).decode("utf-8")

        for idx in range(2, 36):
            emailHsMoi = f"hocsinh{idx}@dss.edu.vn"
            tkMoi = phien.query(TaiKhoan).filter_by(email=emailHsMoi).first()
            if not tkMoi:
                tkMoi = TaiKhoan(
                    email=emailHsMoi,
                    matKhauHash=matKhauChungHash,
                    vaiTro="hocSinh",
                    trangThai="kichHoat"
                )
                phien.add(tkMoi)
                phien.flush()

                vungChon = danhSachVung[idx % len(danhSachVung)]
                nganSachChon = float(10000000 + (idx % 6) * 5000000)
                hsMoi = HocSinh(
                    maTaiKhoan=tkMoi.maTaiKhoan,
                    hoTen=f"Học sinh Thử nghiệm {idx}",
                    soDienThoai=f"098{idx:07d}",
                    namDuTuyen=2026,
                    nganSachToiDa=nganSachChon,
                    kyTinhNganSach="nam",
                    diaBanUuTien=vungChon,
                    batBuocNganSach=False,
                    batBuocDiaBan=False,
                    mucTieuNgheNghiep="Phát triển chuyên môn và xây dựng sự nghiệp bền vững.",
                    uuTienCaNhan=f"Ưu tiên môi trường đào tạo chất lượng tại {vungChon}"
                )
                phien.add(hsMoi)
                phien.flush()
                cacHocSinhId.append(hsMoi.maHocSinh)

                import random
                random.seed(idx * 100)
                cacMon = ["Toan", "Van", "NgoaiNgu", "VatLy", "HoaHoc", "SinhHoc", "LichSu", "DiaLy", "GDCD"]
                for mon in cacMon:
                    diem = round(min(10.0, max(4.0, random.gauss(7.2, 1.3))), 1)
                    phien.add(DiemMonHoc(
                        maHocSinh=hsMoi.maHocSinh,
                        tenMon=mon,
                        diemSo=diem,
                        thangDiem=10.0,
                        loaiDiem="thpt",
                        namHocKy="2026"
                    ))

                nganh1 = danhSachNhomNganh[(idx) % len(danhSachNhomNganh)]
                nganh2 = danhSachNhomNganh[(idx + 1) % len(danhSachNhomNganh)]
                phien.add(SoThichHocSinh(maHocSinh=hsMoi.maHocSinh, nhomNganh=nganh1, mucDoThich=5))
                phien.add(SoThichHocSinh(maHocSinh=hsMoi.maHocSinh, nhomNganh=nganh2, mucDoThich=4))

                phien.add(DanhSachNguyenVong(
                    maHocSinh=hsMoi.maHocSinh,
                    tenDanhSach="Danh sách nguyện vọng chính thức",
                    phienBan=1
                ))
            else:
                hsExist = phien.query(HocSinh).filter_by(maTaiKhoan=tkMoi.maTaiKhoan).first()
                if hsExist and hsExist.maHocSinh not in cacHocSinhId:
                    cacHocSinhId.append(hsExist.maHocSinh)

        phien.flush()

        soLuongMau = phien.query(MauDanhGia).count()
        if soLuongMau == 0:
            logging.info("Khởi tạo tập mẫu đánh giá kỹ thuật (duLieuMoPhong=True) để huấn luyện Random Forest...")
            cacPa = phien.query(PhuongAn).limit(100).all()
            if cacPa:
                import random
                random.seed(42)
                danhSachMau = []
                for hsId in cacHocSinhId:
                    for pa in cacPa[:15]:
                        nhan = round(random.uniform(2.0, 4.8), 2)
                        danhSachMau.append(MauDanhGia(
                            maHocSinh=hsId,
                            maPhuongAn=pa.maPhuongAn,
                            nhanPhuHop=nhan,
                            nguonNhan="moPhongKiemThu",
                            duLieuMoPhong=True,
                            thoiDiemGanNhan=datetime.utcnow()
                        ))
                phien.bulk_save_objects(danhSachMau)
                phien.flush()
                logging.info("Đã tạo %d mẫu đánh giá thử nghiệm kỹ thuật.", len(danhSachMau))

                phienBanDl = PhienBanDuLieu(
                    tenPhienBan="v1.0_duLieuTongHopTuyenSinh",
                    moTa="Dữ liệu tổng hợp tuyển sinh 2021-2026 từ Bộ GD&ĐT kèm nhãn đánh giá mô phỏng kiểm thử kỹ thuật.",
                    soLuongBanGhi=len(danhSachMau),
                    duLieuMoPhong=True
                )
                phien.add(phienBanDl)
                phien.flush()

        phien.commit()
    except Exception as e:
        phien.rollback()
        logging.error("Lỗi khi tạo tài khoản mặc định: %s", e)
        raise
    finally:
        dongPhienKetNoi()


if __name__ == "__main__":
    duongDanExcel = CauHinh.DUONG_DAN_DU_LIEU_GOC
    if os.path.exists(duongDanExcel):
        napDuLieuExcel(duongDanExcel, gioiHanBanGhi=3000)
    else:
        logging.warning("Không tìm thấy tệp Excel tại %s, bỏ qua nạp Excel.", duongDanExcel)
    taoTaiKhoanMacDinhVaMauDanhGia()
